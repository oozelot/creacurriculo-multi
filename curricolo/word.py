"""Elementi comuni per comporre i documenti Word (paragrafi e tabelle diretti)."""

from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .testo import ripara_mojibake
from .config import CARTELLA_LEGACY

MODELLO_INTESTAZIONE = CARTELLA_LEGACY / "modelli_doc" / "01modelloprogrammazionedatibase.docx"
# Il modello ha margini asimmetrici: 15 cm lascia spazio sicuro anche in LibreOffice.
LARGHEZZA_TABELLA = Cm(15)
LARGHEZZA_COLONNA_STRETTA = Cm(0.8)


def crea_documento():
    """Crea un documento con l'intestazione istituzionale dei modelli legacy."""
    doc = Document(MODELLO_INTESTAZIONE)
    corpo = doc._element.body
    for elemento in list(corpo):
        if elemento.tag.endswith("}sectPr"):
            continue
        corpo.remove(elemento)
    stile = doc.styles["Normal"]
    stile.font.name = "Times New Roman"
    stile.font.size = Pt(12)
    return doc


def _distanza_da_tabella(paragrafo) -> None:
    precedente = paragrafo._p.getprevious()
    if precedente is not None and precedente.tag.endswith("}tbl"):
        paragrafo.paragraph_format.space_before = Pt(8)


def paragrafo(doc, testo="", *, grassetto=False, centrato=False, corpo=11, sottolineato=False):
    par = doc.add_paragraph()
    _distanza_da_tabella(par)
    if centrato:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tratto = par.add_run(ripara_mojibake(testo))
    tratto.bold = grassetto
    tratto.underline = sottolineato
    tratto.font.size = Pt(corpo)
    return par


def etichetta_valore(doc, etichetta: str, valore: str, corpo=11) -> None:
    par = doc.add_paragraph()
    _distanza_da_tabella(par)
    testa = par.add_run(f"{etichetta} ")
    testa.bold = True
    testa.font.size = Pt(corpo)
    coda = par.add_run(ripara_mojibake(valore or ""))
    coda.font.size = Pt(corpo)


def testo_multiriga(doc, contenuto: str, corpo=10) -> None:
    righe = [r for r in (contenuto or "").splitlines() if r.strip()]
    for riga in righe or [""]:
        paragrafo(doc, riga.strip(), corpo=corpo)


def separatore_blocco(doc) -> None:
    """Riga vuota fra sezioni composte da modelli Word differenti."""
    spaziatore = doc.add_paragraph()
    spaziatore.paragraph_format.space_after = Pt(8)


def tabella(doc, colonne: int, colonne_strette: tuple[int, ...] = ()):
    corpo = doc._element.body
    precedente = corpo[-2] if len(corpo) > 1 else None
    if precedente is not None and (
        precedente.tag.endswith("}tbl")
        or (precedente.tag.endswith("}p") and precedente.text)
    ):
        spaziatore = doc.add_paragraph()
        spaziatore.paragraph_format.space_after = Pt(8)
    nuova = doc.add_table(rows=0, cols=colonne)
    nuova.style = "Table Grid"
    nuova.alignment = WD_TABLE_ALIGNMENT.CENTER
    nuova.autofit = False
    larghezza_colonna = LARGHEZZA_TABELLA // colonne
    for colonna in nuova.columns:
        colonna.width = larghezza_colonna
    nuova._colonne_strette = set(colonne_strette)
    return nuova


def _larghezze_colonne(valori: list[str], colonne_strette: set[int]) -> list:
    strette = {
        indice
        for indice, valore in enumerate(valori)
        if (indice == 0 and valore.strip().isdigit()) or valore.strip().upper() == "X"
    } | colonne_strette
    spazio_testo = LARGHEZZA_TABELLA - LARGHEZZA_COLONNA_STRETTA * len(strette)
    larghezza_testo = spazio_testo // max(1, len(valori) - len(strette))
    return [LARGHEZZA_COLONNA_STRETTA if indice in strette else larghezza_testo
            for indice in range(len(valori))]


def _imposta_griglia(destinazione, larghezze: list) -> None:
    """Aggiorna le colonne XML, usate da LibreOffice per l'ampiezza reale."""
    for colonna, larghezza in zip(destinazione.columns, larghezze):
        colonna.width = larghezza
    for riga in destinazione.rows:
        for cella, larghezza in zip(riga.cells, larghezze):
            cella.width = larghezza


def riga_tabella(destinazione, valori: list[str], *, grassetto=False, corpo=9) -> None:
    celle = destinazione.add_row().cells
    larghezze = _larghezze_colonne(valori, destinazione._colonne_strette)
    _imposta_griglia(destinazione, larghezze)
    for cella, valore, larghezza in zip(celle, valori, larghezze):
        cella.width = larghezza
        cella.text = ""
        par = cella.paragraphs[0]
        for indice, riga in enumerate(ripara_mojibake(valore or "").split("\n")):
            tratto = (par if indice == 0 else cella.add_paragraph()).add_run(riga)
            tratto.bold = grassetto
            tratto.font.size = Pt(corpo)


def riga_unita(destinazione, titolo: str) -> None:
    """Riga a tutta larghezza: le celle vengono unite come nelle tabelle dei modelli."""
    celle = destinazione.add_row().cells
    unita = celle[0]
    for altra in celle[1:]:
        unita = unita.merge(altra)
    unita.text = ""
    tratto = unita.paragraphs[0].add_run(ripara_mojibake(titolo))
    tratto.bold = True
    tratto.font.size = Pt(10)


def blocco_elenco(doc, titolo: str, testi: list[str], codici: list[str], *, solo_codici=False):
    paragrafo(doc, titolo, grassetto=True)
    valorizzati = [
        (testo, codice)
        for testo, codice in zip(testi, codici)
        if (codice if solo_codici else testo or "").strip()
    ]
    destinazione = tabella(doc, 2, colonne_strette=(0,))
    if not valorizzati:
        riga_tabella(destinazione, ["", ""])
        return
    for numero, (testo, codice) in enumerate(valorizzati, start=1):
        if solo_codici:
            contenuto = codice.strip()
        else:
            contenuto = "\n".join(r.strip() for r in testo.splitlines() if r.strip())
        riga_tabella(destinazione, [str(numero), contenuto])
