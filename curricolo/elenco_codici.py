"""Esportazione dell'elenco aggiornato dei codici PECUP."""

from __future__ import annotations

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from .config import RADICE
from .db import connessione

FILE_ELENCO_CODICI = RADICE / "Elenco codici PECUP.docx"
NOMI_TIPO = {
    "abilita": "Abilita'",
    "conoscenza": "Conoscenze",
    "competenza": "Competenze",
}


def aggiorna_elenco_codici() -> None:
    """Crea nella root il riepilogo Word dei codici PECUP presenti nel database."""
    with connessione() as conn:
        righe = conn.execute(
            """
            SELECT v.disciplina, c.nome AS classe, i.nome AS indirizzo,
                   p.tipo, p.codice, p.descrizione
              FROM pecup p
              JOIN pecup_validita v ON v.pecup_id = p.id
              JOIN classi c ON c.numero = v.classe
              JOIN indirizzi i ON i.id = v.indirizzo_id
             ORDER BY v.disciplina, c.numero, i.posizione, p.tipo, p.codice
            """
        ).fetchall()

    doc = Document()
    sezione = doc.sections[0]
    sezione.orientation = WD_ORIENT.LANDSCAPE
    sezione.page_width, sezione.page_height = sezione.page_height, sezione.page_width
    sezione.top_margin = Cm(1.5)
    sezione.bottom_margin = Cm(1.5)
    sezione.left_margin = Cm(1.5)
    sezione.right_margin = Cm(1.5)
    stile = doc.styles["Normal"]
    stile.font.name = "Times New Roman"
    stile.font.size = Pt(9)

    titolo = doc.add_heading("ELENCO CODICI PECUP", level=0)
    titolo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Elenco aggiornato automaticamente dai dati presenti nel programma.")

    tabella = doc.add_table(rows=1, cols=6)
    tabella.style = "Table Grid"
    intestazioni = ("Materia", "Classe", "Indirizzo", "Tipo", "Codice", "Descrizione")
    for cella, testo in zip(tabella.rows[0].cells, intestazioni):
        cella.text = testo
        for tratto in cella.paragraphs[0].runs:
            tratto.bold = True

    for riga in righe:
        celle = tabella.add_row().cells
        valori = (
            riga["disciplina"],
            riga["classe"],
            riga["indirizzo"],
            NOMI_TIPO.get(riga["tipo"], riga["tipo"]),
            riga["codice"],
            riga["descrizione"],
        )
        for cella, valore in zip(celle, valori):
            cella.text = valore

    doc.save(FILE_ELENCO_CODICI)