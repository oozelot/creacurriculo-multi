"""Catalogo e persistenza del piano di educazione civica."""

from __future__ import annotations

import json
from typing import Any

from .config import CARTELLA_LEGACY
from .db import connessione

MACROAREE = (
    "COSTITUZIONE",
    "EDUCAZIONE SOSTENIBILE",
    "CITTADINANZA DIGITALE",
)

ARTICOLAZIONI_AGRARIO = ("PT", "GAT", "ENO")


def quadrimestre_periodo(periodo: str) -> str:
    """Restituisce I o II in base ai mesi compresi nel periodo del modulo."""
    valore = periodo.strip().lower()
    if not valore:
        return ""
    mesi_primo = ("settembre", "ottobre", "novembre", "dicembre", "gennaio")
    mesi_secondo = ("febbraio", "marzo", "aprile", "maggio", "giugno")
    if any(mese in valore for mese in mesi_secondo):
        return "II"
    if any(mese in valore for mese in mesi_primo):
        return "I"
    return ""


def contesto(indirizzo: str) -> tuple[str, str]:
    """Traduce l'indirizzo della programmazione nel corso Admin e nell'articolazione."""
    valore = indirizzo.upper()
    if "CAT" in valore:
        return "CAT", "COMUNE"
    if "GRAFICO" in valore:
        return "GRAFICO", "COMUNE"
    if "G.A.T" in valore:
        return "AGRARIO", "GAT"
    if "P.T." in valore:
        return "AGRARIO", "PT"
    if "ENO" in valore:
        return "AGRARIO", "ENO"
    if "AGRARIO" in valore:
        return "AGRARIO", "COMUNE"
    return valore, "COMUNE"


def _corsi_piano(corso: str) -> tuple[str, ...]:
    return (corso,) if corso != "COMUNE" else ("CAT", "GRAFICO", "AGRARIO")


def _piano_per_docente(
    classe: str, indirizzo: str, disciplina: str
) -> tuple[str, dict[str, Any]] | None:
    corso, articolazione = contesto(indirizzo)
    from .catalogo import nome_disciplina_visualizzato

    nome = nome_disciplina_visualizzato(disciplina).casefold()
    contesti = (articolazione, "COMUNE") if articolazione != "COMUNE" else ("COMUNE",)
    for corso_piano in _corsi_piano(corso):
        for contesto_piano in contesti:
            for candidato in piani_contesto(corso_piano, classe, contesto_piano):
                if nome_disciplina_visualizzato(candidato["disciplina"]).casefold() == nome:
                    return corso_piano, candidato
    return None


def piano_per_docente(classe: str, indirizzo: str, disciplina: str) -> dict[str, Any] | None:
    risultato = _piano_per_docente(classe, indirizzo, disciplina)
    return risultato[1] if risultato else None


def piano_completo(dati: dict[str, Any] | None) -> bool:
    if not dati or not dati.get("voci"):
        return False
    return all(int(voce.get("ore", 0) or 0) >= 33 for voce in dati["voci"])


def piano_disponibile(dati: dict[str, Any] | None) -> bool:
    return bool(dati and any(int(voce.get("ore", 0) or 0) > 0 for voce in dati.get("voci", [])))


def catalogo_voci() -> dict[str, list[str]]:
    """Legge il catalogo persistente, aggiornato eventualmente all'avvio dal file Excel."""
    risultato = {macroarea: [] for macroarea in MACROAREE}
    with connessione() as conn:
        righe = conn.execute(
            "SELECT macroarea, voce FROM educazione_civica_catalogo ORDER BY macroarea, posizione"
        ).fetchall()
    for riga in righe:
        if riga["macroarea"] in risultato:
            risultato[riga["macroarea"]].append(riga["voce"])
    return risultato


def _leggi_catalogo_excel() -> dict[str, list[str]] | None:
    percorso = CARTELLA_LEGACY / "modelli_doc" / "edcivica.xlsx"
    if not percorso.is_file():
        return None
    from openpyxl import load_workbook

    risultato = {macroarea: [] for macroarea in MACROAREE}
    viste = {macroarea: set() for macroarea in MACROAREE}
    proprietari: set[str] = set()
    macroarea: str | None = None
    foglio = load_workbook(percorso, data_only=True, read_only=True).active
    for riga in foglio.iter_rows(min_col=1, max_col=1, values_only=True):
        valore = str(riga[0] or "").strip()
        if not valore:
            continue
        normalizzato = valore.upper()
        if normalizzato in MACROAREE and normalizzato != macroarea:
            macroarea = normalizzato
        elif macroarea is not None:
            chiave = " ".join(valore.casefold().split())
            if chiave not in viste[macroarea] and chiave not in proprietari:
                risultato[macroarea].append(valore)
                viste[macroarea].add(chiave)
                proprietari.add(chiave)
    return risultato


def sincronizza_catalogo_file() -> None:
    """Sostituisce il catalogo DB solo se il file aggiornabile è disponibile."""
    catalogo = _leggi_catalogo_excel()
    if catalogo is None:
        return
    with connessione() as conn:
        conn.execute("DELETE FROM educazione_civica_catalogo")
        conn.executemany(
            "INSERT INTO educazione_civica_catalogo (macroarea, voce, posizione) VALUES (?, ?, ?)",
            [
                (macroarea, voce, posizione)
                for macroarea, voci in catalogo.items()
                for posizione, voce in enumerate(voci, 1)
            ],
        )
        conn.commit()


def normalizza_piano(dati: dict[str, Any] | None) -> dict[str, Any] | None:
    """Allinea ogni voce alla macroarea ufficiale e rimuove duplicati storici."""
    if dati is None:
        return None
    catalogo = catalogo_voci()
    associazioni = {
        " ".join(voce.casefold().split()): macroarea
        for macroarea, voci in catalogo.items()
        for voce in voci
    }
    voci_normalizzate: dict[str, dict[str, Any]] = {}
    for voce in dati.get("voci", []):
        nome = " ".join(str(voce.get("voce", "")).strip().split())
        if not nome:
            continue
        macroarea = associazioni.get(nome.casefold())
        if not macroarea:
            continue
        chiave = nome.casefold()
        ore = int(voce.get("ore", 0) or 0)
        precedente = voci_normalizzate.get(chiave)
        if precedente is None or ore > int(precedente.get("ore", 0) or 0):
            voci_normalizzate[chiave] = {"macroarea": macroarea, "voce": nome, "ore": ore}
    risultato = dict(dati)
    risultato["voci"] = list(voci_normalizzate.values())
    risultato["ore_disciplina"] = sum(int(voce.get("ore", 0) or 0) for voce in risultato["voci"])
    return risultato


def piano(corso: str, classe: str, articolazione: str, disciplina: str) -> dict[str, Any] | None:
    with connessione() as conn:
        riga = conn.execute(
            "SELECT dati FROM educazione_civica_piani WHERE corso = ? AND classe = ? AND articolazione = ? AND disciplina = ?",
            (corso, classe, articolazione, disciplina),
        ).fetchone()
    return normalizza_piano(json.loads(riga["dati"])) if riga else None


def salva_piano(corso: str, classe: str, articolazione: str, disciplina: str, dati: dict[str, Any]) -> None:
    dati = normalizza_piano(dati) or dati
    contenuto = json.dumps(dati, ensure_ascii=False)
    with connessione() as conn:
        conn.execute(
            """
            INSERT INTO educazione_civica_piani
                (corso, classe, articolazione, disciplina, dati, aggiornato_il)
            VALUES (?, ?, ?, ?, ?, datetime('now', 'localtime'))
            ON CONFLICT(corso, classe, articolazione, disciplina) DO UPDATE SET
                dati = excluded.dati,
                aggiornato_il = excluded.aggiornato_il
            """,
            (corso, classe, articolazione, disciplina, contenuto),
        )
        conn.commit()


def elimina_piano(corso: str, classe: str, articolazione: str, disciplina: str) -> None:
    with connessione() as conn:
        conn.execute(
            "DELETE FROM educazione_civica_piani WHERE corso = ? AND classe = ? AND articolazione = ? AND disciplina = ?",
            (corso, classe, articolazione, disciplina),
        )
        conn.commit()


def registra_conferma(classe: str, indirizzo: str, disciplina: str, chiave: str, dati: dict[str, str]) -> bool:
    corso, articolazione = contesto(indirizzo)
    corsi = _corsi_piano(corso)
    salvata = False
    for corso_piano in corsi:
        piano_dati = piano(corso_piano, classe, articolazione, disciplina)
        if piano_dati is None:
            piano_dati = piano(corso_piano, classe, "COMUNE", disciplina)
            articolazione_piano = "COMUNE"
        else:
            articolazione_piano = articolazione
        if not piano_disponibile(piano_dati):
            continue
        conferme = dict(piano_dati.get("conferme", {}))
        conferme[chiave] = dati
        piano_dati["conferme"] = conferme
        salva_piano(corso_piano, classe, articolazione_piano, disciplina, piano_dati)
        salvata = True
    return salvata


def chiave_conferma(cognome: str, nome: str, disciplina: str, numero_modulo: int, indice_ud: int) -> str:
    return f"{cognome} {nome}|{disciplina}|M{numero_modulo}|UD{indice_ud}"


def rimuovi_conferme_programmazione(prog: Any) -> None:
    corso, articolazione = contesto(prog.indirizzo)
    chiavi = set()
    for numero_modulo, modulo in enumerate(prog.moduli, 1):
        for indice_ud, _ in enumerate(modulo.unita, 1):
            chiavi.add(chiave_conferma(prog.cognome, prog.nome, prog.disciplina, numero_modulo, indice_ud))
            chiavi.add(f"{prog.cognome} {prog.nome}|M{numero_modulo}|UD{indice_ud}")
    for corso_piano in _corsi_piano(corso):
        for articolazione_piano in {articolazione, "COMUNE"}:
            piano_dati = piano(corso_piano, prog.classe, articolazione_piano, prog.disciplina)
            if piano_dati is None:
                continue
            conferme = dict(piano_dati.get("conferme", {}))
            nuove_conferme = {chiave: dati for chiave, dati in conferme.items() if chiave not in chiavi}
            if nuove_conferme != conferme:
                piano_dati["conferme"] = nuove_conferme
                salva_piano(corso_piano, prog.classe, articolazione_piano, prog.disciplina, piano_dati)


def sincronizza_conferme_programmazione(prog: Any) -> None:
    rimuovi_conferme_programmazione(prog)
    for numero_modulo, modulo in enumerate(prog.moduli, 1):
        for indice_ud, unita in enumerate(modulo.unita, 1):
            if 2 not in unita.multidisciplinare or not unita.ec_voci or not unita.ec_ore:
                continue
            registra_conferma(
                prog.classe,
                prog.indirizzo,
                prog.disciplina,
                chiave_conferma(prog.cognome, prog.nome, prog.disciplina, numero_modulo, indice_ud),
                {
                    "voci": list(unita.ec_voci),
                    "quadrimestre": unita.ec_quadrimestre,
                    "periodo": unita.ec_periodo,
                    "ore": unita.ec_ore,
                },
            )


def ore_utilizzate_per_voce(dati: dict[str, Any] | None) -> dict[str, int]:
    """Somma le ore confermate per ogni voce nella materia e nella classe."""
    risultato: dict[str, int] = {}
    for conferma in (dati or {}).get("conferme", {}).values():
        try:
            ore = int(conferma.get("ore", 0) or 0)
        except (TypeError, ValueError):
            continue
        for voce in conferma.get("voci", []):
            risultato[voce] = risultato.get(voce, 0) + ore
    return risultato


def piani_contesto(corso: str, classe: str, articolazione: str = "") -> list[dict[str, Any]]:
    with connessione() as conn:
        righe = conn.execute(
            "SELECT disciplina, articolazione, dati FROM educazione_civica_piani WHERE corso = ? AND classe = ? AND articolazione = ? ORDER BY disciplina",
            (corso, classe, articolazione),
        ).fetchall()
    risultato = []
    for riga in righe:
        dati_originali = json.loads(riga["dati"])
        dati_normalizzati = normalizza_piano(dati_originali) or {}
        if dati_normalizzati != dati_originali:
            salva_piano(corso, classe, riga["articolazione"], riga["disciplina"], dati_normalizzati)
        risultato.append({
            "disciplina": riga["disciplina"],
            "articolazione": riga["articolazione"],
            **dati_normalizzati,
        })
    return risultato


def stato_piano(dati: dict[str, Any] | None) -> str:
    if not dati:
        return "inattivo"
    return "verde" if piano_completo(dati) else "rosso"