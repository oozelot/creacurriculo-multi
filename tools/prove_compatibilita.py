"""Prove di compatibilita' con il programma VB6.

1) Rilegge la cartella di lavoro BIANCHI_ALESSANDRA/PRIMA_COMUNE_CHIMICA prodotta dal VB6,
   rigenera il file da inviare e lo confronta riga per riga con quello realmente ricevuto
   in ADMIN/file_ricevuti (sola lettura: nessun file dell'utente viene modificato).
2) Verifica il ciclo completo scrittura/rilettura e la generazione dei quattro documenti
   Word su una cartella di prova, poi la rimuove.
3) Verifica la modalita' ADMIN: lettura di tutti i file ricevuti, archiviazione, griglia di
   completamento e generazione dei documenti d'istituto (poi rimossi).

Uso: python tools/prove_compatibilita.py
"""

from __future__ import annotations

import shutil
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402

from curricolo import (  # noqa: E402
    archivio,
    completamento,
    documenti_istituto,
    esporta_ini,
    esporta_word,
    importa_ini,
    legacy,
    modello,
)
from curricolo.config import ENCODING_INI, RADICE  # noqa: E402

DOCENTE = ("BIANCHI", "ALESSANDRA", "PRIMA", "COMUNE", "CHIMICA")
FILE_RICEVUTO = (
    RADICE / "ADMIN" / "file_ricevuti" / "CHIMI_COM_BIAALE_PRIMA_06_09_23 - Alessandra Olga Bianchi.ini"
)


def _confronta(prog, percorso_atteso: Path, giorno: date) -> int:
    generate = esporta_ini.genera_righe(prog, giorno)
    attese = percorso_atteso.read_text(encoding=ENCODING_INI, errors="replace").splitlines()

    differenze = 0
    for numero, (ottenuta, attesa) in enumerate(zip(generate, attese), start=1):
        if ottenuta.strip() != attesa.strip():
            differenze += 1
            if differenze <= 5:
                print(f"  riga {numero}\n    atteso  : {attesa!r}\n    ottenuto: {ottenuta!r}")
    if len(generate) != len(attese):
        differenze += 1
        print(f"  numero righe diverso: generate {len(generate)}, attese {len(attese)}")

    esito = "identico" if differenze == 0 else f"{differenze} differenze"
    print(f"  {percorso_atteso.name}: {len(attese)} righe -> {esito}")
    return differenze


def confronto_con_file_reali() -> int:
    """Rigenera il file di invio di ogni lavoro presente e lo confronta con quello del VB6."""
    differenze = 0
    confrontati = 0

    for lavoro in legacy.elenca_lavori():
        prog = modello.carica(**lavoro)
        cartella_invio = legacy.cartella_docente(lavoro["cognome"], lavoro["nome"]) / "da_inviare"
        if not cartella_invio.exists():
            continue
        prefisso = esporta_ini.nome_file_invio(prog)[: -len("dd_mm_yy.ini")]
        for percorso in sorted(cartella_invio.glob(f"{prefisso}*.ini")):
            giorno_testo = percorso.stem[-8:]
            giorno = date(
                2000 + int(giorno_testo[6:8]), int(giorno_testo[3:5]), int(giorno_testo[0:2])
            )
            confrontati += 1
            differenze += _confronta(prog, percorso, giorno)

    if FILE_RICEVUTO.exists() and legacy.esiste_lavoro(*DOCENTE):
        confrontati += 1
        differenze += _confronta(modello.carica(*DOCENTE), FILE_RICEVUTO, date(2023, 9, 6))

    print(f"file confrontati: {confrontati}, differenze totali: {differenze}")
    return differenze


def ciclo_scrittura_lettura() -> int:
    prova = ("PROVA", "AUTOMATICA", "PRIMA", "COMUNE", "CHIMICA")
    if legacy.esiste_lavoro(*DOCENTE):
        prog = modello.carica(*DOCENTE)
    else:
        print("[SALTATO] nessun lavoro di partenza da duplicare")
        return 0

    prog.cognome, prog.nome = "PROVA", "AUTOMATICA"
    cartella = legacy.cartella_docente(*prova[:2])
    try:
        modello.salva(prog)
        riletta = modello.carica(*prova)
        atteso = prog.come_dict()
        ottenuto = riletta.come_dict()
        differenze = [c for c in atteso if atteso[c] != ottenuto[c]]
        print(f"ciclo scrittura/rilettura: {len(differenze)} campi diversi {differenze}")

        differenze += _confronta_cartelle(
            legacy.cartella_corso(*DOCENTE), legacy.cartella_corso(*prova)
        )
        percorso = esporta_ini.scrivi_file(prog, date(2023, 9, 6))
        print(f"file di invio scritto in: {percorso.relative_to(RADICE)}")
        differenze += _genera_documenti(prog)
        return len(differenze)
    finally:
        shutil.rmtree(cartella, ignore_errors=True)


def _genera_documenti(prog) -> list[str]:
    """Ogni documento Word deve essere prodotto e rileggibile."""
    problemi = []
    for tipo in esporta_word.DOCUMENTI:
        percorso = esporta_word.genera(prog, tipo, datetime(2023, 9, 6, 10, 30))
        try:
            documento = Document(str(percorso))
            print(
                f"  {tipo}: {len(documento.paragraphs)} paragrafi, "
                f"{len(documento.tables)} tabelle -> {percorso.name}"
            )
        except Exception as errore:  # il file prodotto non e' leggibile
            problemi.append(f"{tipo}: {errore}")
    return problemi


def _confronta_cartelle(originale: Path, riscritta: Path) -> list[str]:
    """I file riscritti devono coincidere con quelli del VB6.

    Il confronto ignora gli spazi ai bordi delle righe: il VB6 li lasciava (Str() antepone
    uno spazio ai numeri) ma applica Trim in lettura, quindi non trasportano informazione.
    """
    diversi = []
    for percorso in sorted(originale.rglob("datispecific*.ini")) + sorted(
        originale.rglob("datibase*.ini")
    ):
        gemello = riscritta / percorso.relative_to(originale)
        if not gemello.exists():
            diversi.append(str(percorso.relative_to(originale)))
            continue
        attese = percorso.read_text(encoding=ENCODING_INI, errors="replace").splitlines()
        ottenute = gemello.read_text(encoding=ENCODING_INI, errors="replace").splitlines()
        if [r.strip() for r in attese] != [r.strip() for r in ottenute]:
            diversi.append(str(percorso.relative_to(originale)))
    esito = "tutti identici" if not diversi else f"{len(diversi)} diversi: {diversi}"
    print(f"riscrittura dei file di lavoro: {esito}")
    return diversi


def prove_admin() -> int:
    """Rilegge tutti i file ricevuti, li archivia e produce i documenti d'istituto."""
    file_ricevuti = archivio.file_ricevuti()
    if not file_ricevuti:
        print("[SALTATO] nessun file in ADMIN/file_ricevuti")
        return 0

    problemi = 0
    illeggibili = []
    for percorso in file_ricevuti:
        try:
            prog = importa_ini.leggi_file(percorso)
        except Exception as errore:  # file non interpretabile
            illeggibili.append(f"{percorso.name}: {errore}")
            continue
        if not (prog.classe and prog.indirizzo and prog.disciplina and prog.moduli):
            illeggibili.append(f"{percorso.name}: dati identificativi incompleti")
    problemi += len(illeggibili)
    for guasto in illeggibili[:5]:
        print(f"  [ERRORE] {guasto}")
    print(f"file ricevuti letti: {len(file_ricevuti)}, non interpretabili: {len(illeggibili)}")

    archivio.importa_tutti()
    record = archivio.elenca()
    print(f"archivio: {len(record)} programmazioni")

    quadro = completamento.griglia()
    print(
        f"griglia completamento: {len(quadro['righe'])} discipline, "
        f"{quadro['ricevute']}/{quadro['attese']} caselle coperte"
    )

    prodotti = []
    competenze = archivio.unita_per_competenza("PRIMA", ["COMUNE"])
    if competenze:
        prodotti.append(
            documenti_istituto.genera_curricolo(
                "PRIMA", ["COMUNE"], sorted(competenze)[:2], giorno=date(2024, 11, 30)
            )
        )
    aree = archivio.aree_multidisciplinari("PRIMA", ["COMUNE"])
    if aree:
        prodotti.append(
            documenti_istituto.genera_multidisciplinare(
                "PRIMA", ["COMUNE"], sorted(aree), date(2024, 11, 30)
            )
        )

    for percorso in prodotti:
        try:
            documento = Document(str(percorso))
            print(
                f"  {percorso.name}: {len(documento.paragraphs)} paragrafi, "
                f"{len(documento.tables)} tabelle"
            )
        except Exception as errore:  # il documento prodotto non e' leggibile
            problemi += 1
            print(f"  [ERRORE] {percorso.name}: {errore}")
        percorso.unlink(missing_ok=True)

    return problemi


def main() -> int:
    errori = confronto_con_file_reali() + ciclo_scrittura_lettura() + prove_admin()
    print("ESITO:", "OK" if errori == 0 else f"{errori} problemi")
    return 1 if errori else 0


if __name__ == "__main__":
    raise SystemExit(main())
