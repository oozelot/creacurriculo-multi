"""Popola il database dei cataloghi leggendo i file INI legacy in do_not_use/.

Rilanciabile: azzera e ricarica le tabelle di catalogo.
Uso: python tools/importa_cataloghi.py
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from curricolo.config import (  # noqa: E402
    CARTELLA_DISCIPLINE,
    CARTELLA_LEGACY,
    CARTELLA_PROGRAMMAZIONE,
    CLASSI,
    COMPETENZE_EUROPEE,
    ENCODING_INI,
    INDIRIZZI,
    SIGLE_INDIRIZZO,
)
from curricolo.db import connessione, crea_schema, sigla_ini_da_sigla  # noqa: E402
from curricolo.testo import ripara_mojibake  # noqa: E402

# Prefisso del file discipline_*.ini -> indirizzo. Le articolazioni esistono solo dal triennio.
FAMIGLIE = {
    "comuni": "COMUNE",
    "cat": "CAT",
    "grafico": "GRAFICO",
    "agrario": "AGRARIO (tutte le articolazioni)",
    "gat": "AGRARIO (g.a.t.)",
    "pt": "AGRARIO (p.t.)",
    "eno": "AGRARIO (eno)",
}

# Sigle tecniche a 3 caratteri usate nei codici PECUP e nei dati applicativi.
SIGLE_DISCIPLINA = {
    "ALTERNATIVA IRC": "ALT", "BIOTECNOLOGIE AGRARIE": "BTC",
    "BIOTECNOLOGIE VITIVINICOLE": "BTV", "CHIMICA": "CHI",
    "COMPLEMENTI DI MATEMATICA": "CMT", "DIRITTO ED ECONOMIA": "DIR",
    "ECONOMIA, ESTIMO, MARKETING E LEGISLAZIONE": "EST", "ENOLOGIA": "ENO",
    "FISICA": "FIS", "GENIO RURALE": "GNR", "GEOGRAFIA": "GEO",
    "GEOPEDOLOGIA ECONOMIA ED ESTIMO": "GEE",
    "GESTIONE DEL CANTIERE E SICUREZZA DELL'AMBIENTE E DEL TERRITORIO": "GCS",
    "GESTIONE DEL CANTIERE E SICUREZZA DELL'AMBIENTE DI LAVORO": "GCS",
    "GESTIONE DELL'AMBIENTE E DEL TERRITORIO": "GAT", "INGLESE": "ING",
    "IRC (RELIGIONE CATTOLICA)": "REL", "ITALIANO": "ITA", "LABORATORI TECNICI": "LTC",
    "MATEMATICA": "MAT", "ORGANIZZAZIONE E GESTIONE DEI PROCESSI PRODUTTIVI": "OGP",
    "PRODUZIONI ANIMALI": "PAN", "PRODUZIONI VEGETALI": "PVG",
    "PROGETTAZIONE COSTRUZIONI E IMPIANTI": "PCI", "PROGETTAZIONE MULTIMEDIALE": "PRM",
    "S.T.A. (SCIENZE E TECNOLOGIE APPLICATE)": "STA", "SCIENZA DELLA TERRA E BIOLOGIA": "SCI",
    "SCIENZE MOTORIE": "SCM", "STORIA": "STO",
    "T.T.R.G. (TECNOLOGIE E TECNICHE DI RAPPRESENTAZIONE GRAFICA)": "TRG",
    "TECNOLOGIE DEI PROCESSI DI PRODUZIONE": "TPP", "TECNOLOGIE INFORMATICHE": "INF",
    "TEORIA DELLA COMUNICAZIONE": "TCM", "TOPOGRAFIA": "TPG", "TRASFORMAZIONE DEI PRODOTTI": "TRP",
    "VITICOLTURA E DIFESA DELLA VITE": "VIT",
}

SUFFISSI_VARIANTE = (
    " (SOLO SE SPECIFICA PER ARTICOLAZIONE)",
    " (SOLO SE SPECIFICA PER INDIRIZZO)",
)

TIPI_PECUP = {
    "abilita": ("abilita.ini", "estraicodiceabilita.ini"),
    "conoscenza": ("conoscenze.ini", "estraicodiceconoscenze.ini"),
    "competenza": ("competenze.ini", "estraicodicecompetenze.ini"),
}

CODICE_JOLLY = "XXXXXXXXXXXX"

# Diciture estese usate nelle tabelle dei modelli Word, nello stesso ordine dei cataloghi.
TESTI_DOCUMENTO = {
    "strategia": [
        "LEZIONE FRONTALE",
        "LAVORO INDIVIDUALE",
        "LAVORO DI GRUPPO",
        "COOPERATIVE LEARNING",
        "PROCEDURE DI RICERCA",
        "ATTIVITA' LABORATORIALE",
        "BRAIN STORMING",
        "CONVERSAZIONE GUIDATA",
        "PEER TO PEER",
        "PROBLEM SOLVING",
    ],
    "mezzo": [
        "LIBRI DI TESTO",
        "TESTI DI CONSULTAZIONE",
        "SCHEDE PREDISPOSTE",
        "ATTREZZATURE E O STRUMENTI TECNICI",
        "SITO DEL DOCENTE",
    ],
    "strumento": [
        "TEST D'INGRESSO",
        "PROVE INTERDISCIPLINARI",
        "VERIFICHE ALLA FINE DELLE UNITA' DI APPRENDIMENTO",
        "PROVE DISCIPLINARI",
        "PROVE DI COMPETENZA",
        "PRODOTTI INDIVIDUALI DEGLI STUDENTI",
        "PRODOTTI DI GRUPPO DEGLI STUDENTI",
        "PROVE LABORATORIALI",
    ],
}

TITOLI_COMPETENZE_UE = [
    "COMUNICARE IN MADRELINGUA",
    "COMUNICARE IN LINGUE STRANIERE",
    "COMPETENZA MATEMATICA E COMPETENZE DI BASE IN SCIENZA E TECNOLOGIA",
    "COMPETENZE DIGITALI",
    "IMPARARE AD IMPARARE",
    "COMPETENZE INTERPERSONALI, INTERCULTURALI E SOCIALI E COMPETENZA CIVICA",
    "SPIRITO DI INIZIATIVA ED IMPRENDITORIALIT\u00c0",
    "CONSAPEVOLEZZA ED ESPRESSIONE CULTURALE",
]

# Nel VB6 ogni competenza occupava un blocco di 6 righe (A:C) distanziato di 10 righe.
FOGLIO_COMPETENZE = CARTELLA_LEGACY / "modelli_doc" / "competenze.xlsx"
PASSO_BLOCCO = 10
RIGHE_BLOCCO = 6



def righe_utili(percorso: Path) -> list[str]:
    # Il VB6 scartava le righe di lunghezza <= 3: manteniamo lo stesso criterio.
    testo = percorso.read_text(encoding=ENCODING_INI, errors="replace")
    return [ripara_mojibake(r.strip()) for r in testo.splitlines() if len(r.strip()) > 3]


def sigla_disciplina(nome: str) -> str:
    base = nome.upper()
    for suffisso in SUFFISSI_VARIANTE:
        if base.endswith(suffisso):
            base = base[: -len(suffisso)].strip()
    return SIGLE_DISCIPLINA.get(base, base[:5])


def importa_anagrafiche(conn) -> None:
    conn.executemany(
        "INSERT INTO indirizzi (nome, sigla, posizione) VALUES (?, ?, ?)",
        [(nome, SIGLE_INDIRIZZO[nome], i) for i, nome in enumerate(INDIRIZZI, start=1)],
    )
    conn.executemany(
        "INSERT INTO classi (numero, nome) VALUES (?, ?)",
        [(i, nome) for i, nome in enumerate(CLASSI, start=1)],
    )


def importa_opzioni(conn) -> None:
    sorgenti = {
        "strategia": "strategie.ini",
        "mezzo": "mezzi.ini",
        "strumento": "strumenti.ini",
    }
    for gruppo, nome_file in sorgenti.items():
        percorso = CARTELLA_PROGRAMMAZIONE / nome_file
        if not percorso.exists():
            print(f"[ATTENZIONE] manca {percorso}")
            continue
        voci = [v for v in righe_utili(percorso) if v.upper() != "ALTRO..."]
        estesi = TESTI_DOCUMENTO.get(gruppo, [])
        conn.executemany(
            "INSERT INTO opzioni (gruppo, posizione, testo, testo_documento) VALUES (?, ?, ?, ?)",
            [
                (gruppo, i, testo, estesi[i - 1] if i <= len(estesi) else testo)
                for i, testo in enumerate(voci, start=1)
            ],
        )
        print(f"{gruppo}: {len(voci)} voci")


def importa_competenze_ue(conn) -> None:
    if not FOGLIO_COMPETENZE.exists():
        print(f"[ATTENZIONE] manca {FOGLIO_COMPETENZE}")
        return
    from openpyxl import load_workbook

    foglio = load_workbook(FOGLIO_COMPETENZE, data_only=True)["Foglio1"]

    def cella(riga: int, colonna: int) -> str:
        valore = foglio.cell(row=riga, column=colonna).value
        return str(valore).replace("\xa0", " ").strip() if valore is not None else ""

    livelli: list[tuple[int, int, str, str]] = []
    for posizione in range(1, 9):
        prima_riga = 1 + (posizione - 1) * PASSO_BLOCCO
        conn.execute(
            "INSERT INTO competenze_ue (posizione, etichetta, titolo_documento, titolo_esteso,"
            " descrizione) VALUES (?, ?, ?, ?, ?)",
            (
                posizione,
                COMPETENZE_EUROPEE[posizione - 1],
                TITOLI_COMPETENZE_UE[posizione - 1],
                cella(prima_riga, 1),
                cella(prima_riga + 1, 1),
            ),
        )
        for ordine in range(1, RIGHE_BLOCCO - 1):
            riga = prima_riga + 1 + ordine
            etichetta = " ".join(cella(riga, 2).split())
            descrizione = cella(riga, 3)
            if etichetta or descrizione:
                livelli.append((posizione, ordine, etichetta, descrizione))

    conn.executemany(
        "INSERT INTO livelli_ue (posizione, ordine, etichetta, descrizione) VALUES (?, ?, ?, ?)",
        livelli,
    )
    print(f"competenze europee: 8 schede, {len(livelli)} livelli")


def importa_discipline(conn) -> None:
    id_indirizzo = {r["nome"]: r["id"] for r in conn.execute("SELECT id, nome FROM indirizzi")}
    ids: dict[str, int] = {}
    coppie: set[tuple[int, int, int]] = set()

    for percorso in sorted(CARTELLA_DISCIPLINE.glob("discipline_*.ini")):
        gambo = percorso.stem[len("discipline_") :]
        famiglia, classe = gambo[:-1], int(gambo[-1])
        indirizzo = FAMIGLIE.get(famiglia)
        if indirizzo is None:
            print(f"[ATTENZIONE] famiglia sconosciuta: {percorso.name}")
            continue

        for nome in righe_utili(percorso):
            if nome not in ids:
                cur = conn.execute(
                    "INSERT INTO discipline (nome, sigla, sigla_ini) VALUES (?, ?, ?)",
                    (nome, sigla_disciplina(nome), sigla_ini_da_sigla(sigla_disciplina(nome), nome)),
                )
                ids[nome] = int(cur.lastrowid)
            coppie.add((ids[nome], id_indirizzo[indirizzo], classe))

    # Le discipline comuni valgono per tutti gli indirizzi della stessa classe.
    id_comune = id_indirizzo["COMUNE"]
    comuni = {(d, c) for d, i, c in coppie if i == id_comune}
    for disciplina_id, classe in comuni:
        for indirizzo_id in id_indirizzo.values():
            coppie.add((disciplina_id, indirizzo_id, classe))

    conn.executemany(
        "INSERT OR IGNORE INTO offerta_formativa (disciplina_id, indirizzo_id, classe) VALUES (?, ?, ?)",
        sorted(coppie),
    )
    print(f"discipline: {len(ids)} - abbinamenti: {len(coppie)}")


def importa_pecup(conn) -> None:
    id_indirizzo = {
        r["posizione"]: r["id"] for r in conn.execute("SELECT id, posizione FROM indirizzi")
    }

    for tipo, (file_testi, file_codici) in TIPI_PECUP.items():
        percorso_testi = CARTELLA_DISCIPLINE / file_testi
        percorso_codici = CARTELLA_DISCIPLINE / file_codici
        if not percorso_testi.exists() or not percorso_codici.exists():
            print(f"[ATTENZIONE] mancano i file per {tipo}")
            continue

        descrizioni: dict[str, str] = {}
        for riga in righe_utili(percorso_testi):
            codice, descrizione = riga[:12].strip(), riga[12:].strip()
            if codice and codice != CODICE_JOLLY and descrizione:
                descrizioni[codice] = descrizione

        ids: dict[str, int] = {}
        validita: set[tuple[int, str, int, int]] = set()
        orfani = 0

        for riga in righe_utili(percorso_codici):
            codice = riga[:12].strip()
            anni = riga[12:17]
            disciplina = riga[52:].strip()
            if not codice or len(riga) < 53 or not disciplina:
                continue
            if codice not in descrizioni:
                orfani += 1
                continue
            if codice not in ids:
                cur = conn.execute(
                    "INSERT OR IGNORE INTO pecup (tipo, codice, descrizione) VALUES (?, ?, ?)",
                    (tipo, codice, descrizioni[codice]),
                )
                if cur.lastrowid:
                    ids[codice] = int(cur.lastrowid)
                else:
                    riga_db = conn.execute(
                        "SELECT id FROM pecup WHERE tipo = ? AND codice = ?", (tipo, codice)
                    ).fetchone()
                    ids[codice] = int(riga_db["id"])

            # 12 caratteri di codice, 5 bit di classe, poi 5 maschere di 7 bit (una per classe).
            for classe in range(1, 6):
                if anni[classe - 1 : classe] != "1":
                    continue
                inizio = 17 + 7 * (classe - 1)
                maschera = riga[inizio : inizio + 7]
                for posizione in range(1, 8):
                    if maschera[posizione - 1 : posizione] == "1":
                        validita.add(
                            (ids[codice], disciplina.upper(), classe, id_indirizzo[posizione])
                        )

        conn.executemany(
            "INSERT OR IGNORE INTO pecup_validita (pecup_id, disciplina, classe, indirizzo_id)"
            " VALUES (?, ?, ?, ?)",
            sorted(validita),
        )
        print(f"{tipo}: {len(ids)} codici, {len(validita)} validita' (orfani: {orfani})")


def main() -> int:
    conn = connessione()
    crea_schema(conn)
    for tabella in (
        "pecup_validita",
        "pecup",
        "offerta_formativa",
        "discipline",
        "opzioni",
        "livelli_ue",
        "competenze_ue",
        "classi",
        "indirizzi",
    ):
        conn.execute(f"DELETE FROM {tabella}")
    importa_anagrafiche(conn)
    importa_opzioni(conn)
    importa_competenze_ue(conn)
    importa_discipline(conn)
    importa_pecup(conn)
    conn.commit()
    conn.close()
    print("Import completato.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
