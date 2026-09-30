from datetime import date

import pytest
from docx import Document
from flask import jsonify

from curricolo import catalogo, documenti_istituto
import main

app = main.app


def test_documento_educazione_civica_riporta_totali_e_nota_irc(monkeypatch, tmp_path):
    monkeypatch.setattr(documenti_istituto, "CARTELLA_ADMIN_OUTPUT", tmp_path)
    monkeypatch.setattr(documenti_istituto, "crea_documento", Document)

    def discipline(classe, indirizzi):
        if classe == "TERZA" and "AGRARIO (tutte le articolazioni)" in indirizzi:
            return ["ALTERNATIVA IRC", "IRC (RELIGIONE CATTOLICA)"]
        return []

    def piani(corso, classe, articolazione):
        if corso != "AGRARIO" or classe != "TERZA" or articolazione != "COMUNE":
            return []
        return [
            {
                "disciplina": nome,
                "voci": [{"macroarea": "COSTITUZIONE", "voce": "LEGALITA", "ore": 1}],
                "ore_disciplina": 1,
                "conferme": {},
            }
            for nome in ("ALTERNATIVA IRC", "IRC (RELIGIONE CATTOLICA)")
        ]

    monkeypatch.setattr(catalogo, "discipline_normalizzate", discipline)
    monkeypatch.setattr(documenti_istituto.educazione_civica, "piani_contesto", piani)
    completamento = {
        (corso, classe): 35
        for corso in ("CAT", "GRAFICO", "AGRARIO")
        for classe in ("PRIMA", "SECONDA", "TERZA", "QUARTA", "QUINTA")
    }
    completamento[("AGRARIO", "TERZA")] = {
        "articolazioni": [
            {"nome": "PT", "ore": 35},
            {"nome": "GAT", "ore": 34},
            {"nome": "ENO", "ore": 36},
        ],
        "stato": "verde",
    }

    percorso = documenti_istituto.genera_educazione_civica(
        "AGRARIO",
        "TERZA",
        completamento=completamento,
        giorno=date(2026, 9, 30),
    )
    documento = Document(percorso)

    testo_tabelle = "\n".join(
        cella.text
        for tabella_piano in documento.tables
        for riga in tabella_piano.rows
        for cella in riga.cells
    )
    assert "TOTALE ORE EDUCAZIONE CIVICA: PT: 35 - GAT: 34 - ENO: 36" in testo_tabelle
    assert "IRC (RELIGIONE CATTOLICA)*" in testo_tabelle
    assert "ALTERNATIVA IRC*" in testo_tabelle
    assert sum(
        paragrafo.text
        == "(*) esclusivamente per gli studenti che si avvalgono di tale materia."
        for paragrafo in documento.paragraphs
    ) == 1


@pytest.mark.parametrize(("formato", "pdf"), [("word", False), ("pdf", True)])
def test_creazione_documento_passa_i_totali_anche_per_pdf(
    monkeypatch, tmp_path, formato, pdf
):
    completamento = [
        {
            "corso": "AGRARIO",
            "ore": [
                35,
                35,
                {
                    "articolazioni": [
                        {"nome": "PT", "ore": 35},
                        {"nome": "GAT", "ore": 34},
                        {"nome": "ENO", "ore": 36},
                    ],
                    "stato": "verde",
                },
                35,
                35,
            ],
        },
    ]
    ricevuti = {}
    percorso_word = tmp_path / "programma.docx"
    percorso_pdf = tmp_path / "programma.pdf"
    monkeypatch.setattr(main, "_calcola_completamento", lambda: completamento)

    def genera(corso, classe, *, completamento, avanzamento):
        ricevuti["contesto"] = completamento[(corso, classe)]
        return percorso_word

    def converti(percorso):
        ricevuti["converti_pdf"] = percorso
        return percorso_pdf

    def avvia(lavoro, **_):
        risultato = lavoro(lambda _messaggio: None)
        return jsonify(risultato)

    monkeypatch.setattr(documenti_istituto, "genera_educazione_civica", genera)
    monkeypatch.setattr(documenti_istituto, "converti_pdf", converti)
    monkeypatch.setattr(main, "_avvia_creazione", avvia)

    risposta = app.test_client().post(
        f"/admin/educazione-civica/crea?formato={formato}",
        data={"corso": "AGRARIO", "classe": "TERZA"},
    )

    assert risposta.status_code == 200
    assert ricevuti["contesto"] == completamento[0]["ore"][2]
    assert ("converti_pdf" in ricevuti) is pdf
    if pdf:
        assert risposta.get_json()["percorso"] == str(percorso_pdf)
